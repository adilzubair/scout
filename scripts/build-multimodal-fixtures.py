"""Create an authored content smoke corpus in an ignored directory, never personal files."""

import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import subprocess
import sys
import wave

ROOT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    root = args.output.resolve()
    if root.exists():
        raise FileExistsError('Refusing to overwrite a fixture directory')
    root.mkdir(parents=True)
    from PIL import Image, ImageDraw
    import av
    import numpy as np
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from openpyxl import Workbook
    from reportlab.pdfgen import canvas

    items = []
    queries = []
    def item(identifier, modality, filename, query, **extra):
        items.append(dict(id=identifier, modality=modality, file=filename, **extra))
        queries.append(dict(id=identifier, query=query, expected=identifier, family=modality))

    for key, text, query in [
        ('pdf-invoice', 'Invoice number 714. Customer Cedar Labs. Total due 2450 dollars. Payment is overdue.', 'the overdue customer invoice for Cedar Labs'),
        ('pdf-menu', 'Dinner menu. Vegetable soup, roasted potatoes and lemon cake.', 'a dinner menu with vegetable soup and lemon cake')]:
        filename = key + '.pdf'
        document = canvas.Canvas(str(root / filename))
        document.drawString(30, 800, text)
        document.save()
        item(key, 'document', filename, query)

    document = Document()
    document.add_heading('Project meeting', 0)
    document.add_paragraph('The deployment was postponed because the security review found an expired certificate.')
    document.save(root / 'meeting.docx')
    item('docx-meeting', 'document', 'meeting.docx', 'meeting notes about postponing deployment after a security review')
    book = Workbook()
    sheet = book.active
    sheet.title = 'Quarterly sales'
    sheet.append(['Region', 'Revenue', 'Quarter'])
    sheet.append(['North', 42000, 'Q3'])
    sheet.append(['South', 17000, 'Q3'])
    book.save(root / 'sales.xlsx')
    item('xlsx-sales', 'spreadsheet', 'sales.xlsx', 'quarterly revenue of forty two thousand in the North region')

    chart = Image.new('RGB', (640, 480), 'white')
    draw = ImageDraw.Draw(chart)
    draw.text((50, 20), 'QUARTERLY REVENUE', fill='black', font_size=30)
    for i, height in enumerate([80, 150, 250, 350]):
        draw.rectangle((80 + 120*i, 420-height, 160 + 120*i, 420), fill='blue')
        draw.text((100 + 120*i, 440), f'Q{i+1}', fill='black', font_size=20)
    chart.save(root / 'chart.png')
    item('image-chart', 'image', 'chart.png', 'a blue bar chart showing revenue increasing every quarter')
    error = Image.new('RGB', (640, 480), 'white')
    draw = ImageDraw.Draw(error)
    draw.rectangle((15, 15, 625, 70), fill='navy')
    draw.text((40, 28), 'CHECKOUT', fill='white', font_size=30)
    draw.text((40, 160), 'PAYMENT FAILED', fill='red', font_size=40)
    draw.text((40, 235), 'Card declined. Try another payment method.', fill='black', font_size=24)
    error.save(root / 'checkout.png')
    item('image-payment', 'image', 'checkout.png', 'a screenshot of a failed payment with a declined card')

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[5])
    slide.shapes.title.text = 'Quarterly revenue growth'
    slide.shapes.add_picture(str(root / 'chart.png'), Inches(1), Inches(1.5), width=Inches(7))
    slides.save(root / 'growth.pptx')
    item('pptx-chart', 'presentation', 'growth.pptx', 'presentation showing quarterly revenue growth in a blue bar chart')
    # A scanned document has no text layer and must exercise visual indexing.
    error.save(root / 'scan.pdf', 'PDF')
    item('pdf-scan', 'visual-document', 'scan.pdf', 'a scanned document showing a payment failure')

    for key, speech, query in [
        ('audio-deployment', 'The project deployment is postponed until the security review is complete.', 'someone says the deployment is postponed until the security review is complete'),
        ('audio-cooking', 'To prepare dinner, chop the vegetables and cook the soup for twenty minutes.', 'a person explaining how to chop vegetables and cook soup')]:
        env = os.environ.copy()
        env['SCOUT_FIXTURE_AUDIO_PATH'] = str(root / (key + '.wav'))
        env['SCOUT_FIXTURE_TEXT'] = speech
        script = 'Add-Type -AssemblyName System.Speech; $v = New-Object System.Speech.Synthesis.SpeechSynthesizer; try { $v.SetOutputToWaveFile($env:SCOUT_FIXTURE_AUDIO_PATH); $v.Speak($env:SCOUT_FIXTURE_TEXT) } finally { $v.Dispose() }'
        subprocess.run(['powershell.exe', '-NoProfile', '-NonInteractive', '-Command', script], env=env,
                       check=True, timeout=30, creationflags=subprocess.CREATE_NO_WINDOW)
        item(key, 'audio', key + '.wav', query)
    with wave.open(str(root / 'beeps.wav'), 'wb') as audio:
        audio.setparams((1, 2, 16000, 0, 'NONE', 'not compressed'))
        samples = [int(16000 * math.sin(2 * math.pi * 1000 * n / 16000)) if n % 16000 < 4000 else 0 for n in range(48000)]
        audio.writeframes(np.asarray(samples, dtype='<i2').tobytes())
    item('audio-beeps', 'audio', 'beeps.wav', 'short electronic beeps separated by silence')

    for key, color, direction, query in [
        ('video-red', 'red', 1, 'a red circle moving from left to right'),
        ('video-blue', 'blue', -1, 'a blue circle moving from right to left')]:
        with av.open(str(root / (key + '.mp4')), 'w') as container:
            stream = container.add_stream('mpeg4', rate=10)
            stream.width, stream.height, stream.pix_fmt = 320, 240, 'yuv420p'
            for number in range(30):
                image = Image.new('RGB', (320, 240), 'white')
                x = 20 + number * 8 if direction == 1 else 260 - number * 8
                ImageDraw.Draw(image).ellipse((x, 90, x+40, 130), fill=color)
                frame = av.VideoFrame.from_image(image)
                for packet in stream.encode(frame):
                    container.mux(packet)
            for packet in stream.encode():
                container.mux(packet)
        item(key, 'video', key + '.mp4', query)

    (root / 'retry.py').write_text('def retry(attempts, action):\n    for attempt in range(attempts):\n        try:\n            return action()\n        except TimeoutError:\n            continue\n    raise TimeoutError("Retries exhausted")\n', encoding='utf-8')
    item('code-retry', 'code', 'retry.py', 'a function that retries an operation after a timeout')
    files = {p.name: dict(bytes=p.stat().st_size, sha256=hashlib.sha256(p.read_bytes()).hexdigest())
             for p in root.iterdir() if p.is_file()}
    manifest = dict(schemaVersion=1, purpose='Authored smoke fixtures, not a quality holdout',
        items=items, queries=queries, files=files,
        audioSource='Windows System.Speech synthetic voice; exact bytes are hashed',
        videoSource='Programmatically drawn circles, MPEG-4 encoded with PyAV')
    (root / 'fixtures.json').write_text(json.dumps(manifest, indent=2) + '\n', encoding='utf-8')
    print(f'Created {len(items)} content fixtures and queries')


if __name__ == '__main__':
    main()
