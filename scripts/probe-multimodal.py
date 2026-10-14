"""Measure full local embeddings on authored content. Not a production admission policy."""

import argparse
import importlib.metadata
import json
import os
import platform
from pathlib import Path
import statistics
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'runtime'))
from semantic_search.embedding import FullEmbedder, local_file
from semantic_search.manifest import MANIFEST, sha256


def memory():
    import psutil
    info = psutil.Process().memory_info()
    return {key: getattr(info, field, None) for key, field in
            [('workingSetBytes', 'rss'), ('peakWorkingSetBytes', 'peak_wset'), ('privateBytes', 'private')]}


def fixture_content(item, root):
    """Small fixture readers only; production segmentation is a later milestone."""
    path = local_file(root / item['file'], root)
    if path.stat().st_size > 64 * 1024 * 1024:
        raise ValueError('Fixture exceeds 64 MiB input budget')
    kind = item['modality']
    if kind == 'code':
        return path.read_text(encoding='utf-8'), 'source-text'
    if kind == 'document':
        if path.suffix == '.pdf':
            from pypdf import PdfReader
            return '\n'.join(page.extract_text() for page in PdfReader(path).pages), 'pdf-text'
        from docx import Document
        return '\n'.join(p.text for p in Document(path).paragraphs), 'docx-paragraphs'
    if kind == 'spreadsheet':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return '\n'.join(sheet.title + '\n' + '\n'.join(' | '.join(str(c) for c in row) for row in
                        sheet.iter_rows(values_only=True)) for sheet in book), 'xlsx-values'
        finally:
            book.close()
    if kind == 'image':
        from PIL import Image
        with Image.open(path) as image:
            if image.width * image.height > 4_000_000:
                raise ValueError('Fixture image exceeds pixel budget')
            return {'image': image.convert('RGB')}, 'image-pixels'
    if kind == 'visual-document':
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(path) as document:
            page = document[0]
            bitmap = page.render(scale=1)
            try:
                return {'image': bitmap.to_pil().convert('RGB')}, 'pdf-rendered-page'
            finally:
                bitmap.close()
                page.close()
    if kind == 'presentation':
        from pptx import Presentation
        from PIL import Image
        import io
        for slide in Presentation(path).slides:
            for shape in slide.shapes:
                if hasattr(shape, 'image'):
                    with Image.open(io.BytesIO(shape.image.blob)) as image:
                        return {'image': image.convert('RGB')}, 'pptx-embedded-image'
        raise ValueError('Fixture lacks an embedded image; complete Office rendering is a separate gate')
    if kind == 'audio':
        import soundfile as sf
        import torch
        import torchaudio.functional as audio
        with sf.SoundFile(path) as stream:
            if len(stream) / stream.samplerate > 60:
                raise ValueError('Audio fixture exceeds 60 seconds; segmentation required')
            samples = stream.read(dtype='float32', always_2d=True).mean(axis=1)
            samples = audio.resample(torch.from_numpy(samples), stream.samplerate, 16000).numpy()
        return {'audio': samples}, 'audio-mono-16khz'
    if kind == 'video':
        import av
        import numpy as np
        with av.open(str(path)) as container:
            if container.duration is None or container.duration / av.time_base > 10:
                raise ValueError('Video fixture exceeds ten seconds or lacks duration')
            stream = container.streams.video[0]
            fps = float(stream.average_rate or 0)
            if fps <= 0 or stream.width * stream.height > 4_000_000:
                raise ValueError('Invalid frame rate or excessive video resolution')
            frames, indices, next_second = [], [], 0
            for index, frame in enumerate(container.decode(stream)):
                if index > 1000 or frame.time is None or frame.time > 10:
                    raise ValueError('Video decoding budget exceeded')
                if frame.time >= next_second:
                    if frame.width * frame.height > 4_000_000:
                        raise ValueError('Decoded frame exceeds pixel budget')
                    frames.append(frame.to_ndarray(format='rgb24'))
                    indices.append(index)
                    next_second += 1
                    if len(frames) > 16:
                        raise ValueError('Video frame budget exceeded')
            if not frames:
                raise ValueError('No video frames decoded')
            metadata = dict(total_num_frames=index+1, fps=fps, duration=container.duration / av.time_base,
                            frames_indices=indices, video_backend='scout-pyav')
        return ({'video': np.stack(frames)}, 'video-pyav-one-fps',
                {'video': {'do_sample_frames': False, 'video_metadata': metadata}})
    raise ValueError('Unsupported fixture modality')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--fixtures', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--model', type=Path, default=ROOT / '.runtime/semantic-search/model')
    parser.add_argument('--device', choices=['cpu', 'cuda'], default='cpu')
    parser.add_argument('--threads', type=int, default=4)
    parser.add_argument('--dimensions', type=int, choices=[768, 512, 256, 128], default=768)
    args = parser.parse_args()
    if args.output.exists():
        raise FileExistsError('Refusing to overwrite a probe result')
    fixture = json.loads(args.fixtures.read_text(encoding='utf-8'))
    required = {'document', 'spreadsheet', 'presentation', 'image', 'visual-document', 'audio', 'video', 'code'}
    if fixture.get('schemaVersion') != 1 or not required.issubset({i['modality'] for i in fixture['items']}):
        raise ValueError('A full content fixture pack is required')
    ids = [i['id'] for i in fixture['items']]
    if len(set(ids)) != len(ids) or any(c['expected'] not in ids for c in fixture['queries']):
        raise ValueError('Duplicate fixture IDs or unknown expected item')
    root = args.fixtures.resolve().parent
    for name, spec in fixture['files'].items():
        path = local_file(root / name, root)
        if path.stat().st_size != spec['bytes'] or sha256(path) != spec['sha256']:
            raise ValueError('Fixture checksum mismatch')
    started = time.perf_counter()
    model = FullEmbedder(args.model, args.device, args.threads, args.dimensions)
    cold_seconds = time.perf_counter() - started
    import numpy as np
    import torch
    import av
    import psutil

    def measured(content, **kwargs):
        if args.device == 'cuda':
            torch.cuda.synchronize()
        start = time.perf_counter()
        vector, tokens = model.embed(content, **kwargs)
        if args.device == 'cuda':
            torch.cuda.synchronize()
        return vector, tokens, (time.perf_counter() - start) * 1000

    rows, vectors = [], {}
    for item in fixture['items']:
        try:
            start = time.perf_counter()
            prepared = fixture_content(item, root)
            content, method = prepared[:2]
            processing = prepared[2] if len(prepared) == 3 else None
            extraction_ms = (time.perf_counter() - start) * 1000
            vector, tokens, elapsed = measured(content, processing_kwargs=processing)
            vectors[item['id']] = vector
            rows.append(dict(id=item['id'], modality=item['modality'], reader=method,
                extractionMs=extraction_ms, embeddingMs=elapsed, tokens=tokens, status='ok'))
            print(f'{item["id"]}: {elapsed:.1f} ms, {tokens} tokens', flush=True)
        except Exception as error:
            rows.append(dict(id=item['id'], modality=item['modality'], status='failed', error=str(error)))
            print(f'{item["id"]}: failed ({type(error).__name__})', flush=True)
    answers = []
    for case in fixture['queries']:
        try:
            vector, tokens, elapsed = measured(case['query'], query=True)
            ranked = sorted(((identifier, float(np.dot(vector, candidate))) for identifier, candidate in
                             vectors.items()), key=lambda pair: (-pair[1], pair[0]))
            answers.append(dict(id=case['id'], family=case['family'], expected=case['expected'],
                top=ranked[:5], correct=bool(ranked and ranked[0][0] == case['expected']),
                embeddingMs=elapsed, tokens=tokens))
        except Exception as error:
            answers.append(dict(id=case['id'], status='failed', error=str(error), correct=False))
    packages = ['sentence-transformers', 'transformers', 'torch', 'torchvision', 'torchaudio', 'av',
                'numpy', 'Pillow', 'soundfile', 'pypdf', 'pypdfium2']
    report = dict(schemaVersion=1, purpose='Full-modality runtime smoke screen, not Scout quality or abstention',
        modelRevision=model.manifest['revision'], modelManifestSha256=sha256(MANIFEST),
        probeSha256=sha256(Path(__file__)),
        embedderSha256=sha256(ROOT / 'runtime/semantic_search/embedding.py'),
        fixtureSha256=sha256(args.fixtures), device=args.device, dtype=model.dtype, threads=args.threads,
        dimensions=args.dimensions, totalColdSeconds=cold_seconds, integritySeconds=model.integrity_seconds,
        hardware=dict(platform=platform.platform(), processor=platform.processor(),
            physicalCores=psutil.cpu_count(logical=False), logicalCores=psutil.cpu_count(),
            physicalMemoryBytes=psutil.virtual_memory().total,
            gpu=torch.cuda.get_device_name() if args.device == 'cuda' else None),
        modelLoadSeconds=model.load_seconds, memory=memory(),
        cudaPeakAllocatedBytes=torch.cuda.max_memory_allocated() if args.device == 'cuda' else None,
        dependencies={p: importlib.metadata.version(p) for p in packages},
        ffmpegLibraries={k: list(v) for k, v in av.library_versions.items()},
        ffmpegBuild=av._core.library_meta,
        inputs=rows, answers=answers, correct=sum(a['correct'] for a in answers), cases=len(answers),
        medianQueryEmbeddingMs=statistics.median([a['embeddingMs'] for a in answers if 'embeddingMs' in a]) if any('embeddingMs' in a for a in answers) else None)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open('x', encoding='utf-8') as stream:
        json.dump(report, stream, indent=2, allow_nan=False)
        stream.write('\n')
    print(json.dumps({k: report[k] for k in ['correct', 'cases', 'totalColdSeconds', 'medianQueryEmbeddingMs', 'memory', 'cudaPeakAllocatedBytes']}))
    return int(any(r['status'] == 'failed' for r in rows) or any(a.get('status') == 'failed' for a in answers))


if __name__ == '__main__':
    raise SystemExit(main())
