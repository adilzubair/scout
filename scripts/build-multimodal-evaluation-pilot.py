"""Author a development-only evaluator pilot. Never counted as release holdout."""
import argparse
from fractions import Fraction
import hashlib
import json
import math
import os
from pathlib import Path
import subprocess
import wave

def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--output',type=Path,required=True);args=parser.parse_args()
    root=args.output.absolute()
    if root.exists():raise FileExistsError('Refusing to overwrite an evaluation corpus')
    (root/'sources').mkdir(parents=True)
    from PIL import Image,ImageDraw,ImageFont
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    import av

    sources=[];tasks=[];narrations=[]
    def path(identifier,extension):return root/'sources'/('f-'+hashlib.sha256(identifier.encode()).hexdigest()[:12]+extension)
    def add(identifier,family,group,extension,grounding,license='Original authored Scout development fixture; not a quality holdout'):
        file=path(identifier,extension)
        sources.append(dict(id=identifier,group=group,split='development',family=family,file=file.relative_to(root).as_posix(),grounding=grounding,license=license))
        return file
    def task(identifier,family,group,query,kind,expected):
        tasks.append(dict(id=identifier,group=group,split='development',family=family,query=query,kind=kind,expected=[dict(source=source,locator=locator) for source,locator in expected]))
    def cases(family,group,positive,negative,ambiguous,locators):
        task(family+'-positive',family,group,positive,'positive',[(family+'-a',locators[0])])
        task(family+'-negative',family,group,negative,'negative',[])
        task(family+'-ambiguous',family,group,ambiguous,'ambiguous',[(family+'-a',locators[0]),(family+'-b',locators[1])])

    for tag,body in [('a','Donated books must be sorted on the upstairs tables. The side door is a fire escape and must not be blocked by boxes.'),('b','Volunteers should sort donated books in the basement. The side door can be used for deliveries between ten and eleven.')]:
        file=add('documents-'+tag,'documents','library-rota','.docx',body)
        doc=Document();doc.add_heading('Library volunteer handover',0);doc.add_paragraph(body);doc.add_paragraph('Return the trolley to the storage cupboard before leaving.');doc.save(file)
    cases('documents','library-rota','document explaining why boxes of donated books must stay away from the side door','document saying donated books must be burned after one month','documents about library volunteers sorting donated books',[dict(kind='paragraphs',startLine=2,endLine=2,section='document')]*2)
    for tag,rows in [('a',[['Cactus','Friday','Every other week',200],['Fern','Monday','Weekly',350]]),('b',[['Cactus','Tuesday','Weekly',150],['Fern','Thursday','Twice weekly',400]])]:
        file=add('spreadsheets-'+tag,'spreadsheets','greenhouse-watering','.xlsx','Greenhouse watering schedule: '+repr(rows))
        book=Workbook();sheet=book.active;sheet.title='Watering';sheet.append(['Plant','Day','Frequency','Millilitres']);[sheet.append(row) for row in rows];book.save(file)
    cases('spreadsheets','greenhouse-watering','spreadsheet scheduling cactus watering every other Friday','spreadsheet listing watering instructions for an orchid collection','spreadsheets containing greenhouse watering schedules',[dict(kind='sheet',sheet='Watering',startCell='A2',endCell='D2')]*2)
    for tag,body in [('a','Disconnect the antenna cable before replacing the battery pack. Inspect the connector seal before reconnecting the antenna.'),('b','Leave the antenna connected while charging the battery. Check the channel selector before starting the radio test.')]:
        file=add('presentations-'+tag,'presentations','radio-maintenance','.pptx',body)
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text='Portable radio workshop';slide.placeholders[1].text=body;deck.save(file)
    cases('presentations','radio-maintenance','slides telling technicians to disconnect the antenna before changing a radio battery','slides about replacing a satellite dish on a rooftop','presentations about portable radio maintenance',[dict(kind='slide',page=1)]*2)
    for tag,body in [('a','def bit_parity(value):\n    """Return one for an odd number of set bits in a byte."""\n    return (value & 255).bit_count() & 1\n'),('b','def byte_sum(values):\n    """Add byte values and wrap the result to eight bits."""\n    return sum(values) & 255\n')]:
        file=add('code-'+tag,'code','byte-checksums','.py',body);file.write_text(body,encoding='utf-8')
    cases('code','byte-checksums','code that checks whether a byte has an odd number of set bits','code that encrypts a message using an RSA public key','code implementing a small checksum over byte values',[dict(kind='lines',startLine=1,endLine=3)]*2)
    font=ImageFont.load_default(size=32);small=ImageFont.load_default(size=25)
    for tag,title,body,color in [('a','PLATFORM 8 CLOSED','Signal inspection in progress','#ac1723'),('b','PLATFORM 3 OPEN','Escalator repairs completed','#147040')]:
        file=add('images-'+tag,'images','rail-notices','.png',f'Authored rail service notice: {title}; {body}; header color {color}.')
        image=Image.new('RGB',(800,500),'white');draw=ImageDraw.Draw(image);draw.rectangle((25,25,775,100),fill=color);draw.text((50,45),'COASTLINE RAIL',font=font,fill='white');draw.text((50,170),title,font=font,fill=color);draw.text((50,250),body,font=small,fill='black');image.save(file)
    cases('images','rail-notices','image announcing a railway platform closure during signal inspection','image showing a railway train derailed beside a tunnel','images of railway service notices',[dict(kind='image',page=1)]*2)
    for tag,text in [('a',"For tonight's rehearsal, keep the emergency exit clear. Leave costume cases behind the stage curtain."),('b',"For tonight's rehearsal, string players should enter quietly after the flute solo. Keep the final chord short.")]:
        file=add('audio-'+tag,'audio','rehearsal-recordings','.wav','Synthetic speech says: '+text,'Original authored script; installed Windows synthetic voice, private evaluation only; preserve exact audio bytes.')
        narrations.append(dict(path=str(file),text=text))
    narration_file=root/'narrations.json';narration_file.write_text(json.dumps(narrations),encoding='utf-8')
    command='Add-Type -AssemblyName System.Speech; $v=New-Object System.Speech.Synthesis.SpeechSynthesizer; try { foreach($n in (Get-Content -LiteralPath $env:SCOUT_EVAL_NARRATIONS -Raw | ConvertFrom-Json)) { $v.SetOutputToWaveFile($n.path); $v.Speak($n.text); $v.SetOutputToNull() } } finally { $v.Dispose() }'
    subprocess.run(['powershell.exe','-NoProfile','-NonInteractive','-Command',command],env={**os.environ,'SCOUT_EVAL_NARRATIONS':str(narration_file)},check=True,timeout=90,creationflags=subprocess.CREATE_NO_WINDOW)
    audio_locators=[]
    for narration in narrations:
        with wave.open(narration['path']) as audio:
            duration=audio.getnframes()/audio.getframerate()
            if not 2<duration<30:raise ValueError('Unexpected synthetic narration duration')
            audio_locators.append(dict(kind='time',startSeconds=0,endSeconds=min(5,duration)))
    cases('audio','rehearsal-recordings','audio asking performers not to obstruct the emergency exit','audio announcing that all concert tickets will be refunded','recordings giving rehearsal instructions',audio_locators)
    for tag,bin_color,parcel_color in [('a','#125ac0','#efcc14'),('b','#b12424','#42a043')]:
        file=add('video-'+tag,'video','robot-sorting','.mp4',f'Silent authored animation: a {parcel_color} parcel is moved by a robot arm from a conveyor into a {bin_color} bin, arriving at six seconds. Not a natural video quality sample.')
        with av.open(str(file),'w') as output:
            stream=output.add_stream('libx264',rate=5);stream.width=480;stream.height=320;stream.pix_fmt='yuv420p';stream.options={'preset':'ultrafast','crf':'20'};stream.codec_context.thread_count=1
            for number in range(40):
                t=number/5;image=Image.new('RGB',(480,320),'#eef0f2');draw=ImageDraw.Draw(image);draw.text((20,18),'SORTING CELL',font=small,fill='black');draw.rectangle((25,240,330,275),fill='#64686a');draw.rectangle((350,180,450,275),fill=bin_color)
                x=60+330*min(t/6,1);y=225-90*math.sin(math.pi*min(t/6,1));draw.line((240,85,x,y),fill='#4b4e50',width=13);draw.rectangle((x-17,y-15,x+17,y+15),fill=parcel_color,outline='black',width=2)
                frame=av.VideoFrame.from_image(image);frame.pts=number;frame.time_base=Fraction(1,5)
                for packet in stream.encode(frame):output.mux(packet)
            for packet in stream.encode(None):output.mux(packet)
    cases('video','robot-sorting','video of a robot placing a yellow parcel into a blue bin','video of a forklift carrying wooden pallets outdoors','videos showing a robotic sorting station',[dict(kind='time',startSeconds=4,endSeconds=7)]*2)
    for source in sources:
        file=root/source['file'];source.update(bytes=file.stat().st_size,sha256=hashlib.sha256(file.read_bytes()).hexdigest())
    dataset=dict(schemaVersion=1,id='multimodal-evaluator-pilot-v1',purpose='development-pilot',now='2026-10-15T00:00:00Z',languages=['en'],sources=sources,tasks=tasks)
    (root/'dataset.json').write_text(json.dumps(dataset,indent=2,ensure_ascii=False),encoding='utf-8')
    print(f'Authored {len(sources)} sources and {len(tasks)} development tasks: {root}')
    print('This pilot is not a holdout. Natural photos, environmental recordings, real video, joint AV, larger content and multilingual slices remain required.')

if __name__=='__main__':main()
