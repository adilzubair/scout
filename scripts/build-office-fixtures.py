"""Author small layout fixtures in all six supported Office formats."""

import argparse
import json
from pathlib import Path
import sys
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'runtime'))
from semantic_search.manifest import sha256

NS = '''xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"
xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
xmlns:xlink="http://www.w3.org/1999/xlink"'''


def write_odf(path, kind, body, styles='', assets=None):
    mime = 'application/vnd.oasis.opendocument.' + kind
    entries = {'content.xml': f'''<?xml version="1.0" encoding="UTF-8"?>
<office:document-content {NS} office:version="1.3">
<office:automatic-styles>{styles}</office:automatic-styles>
<office:body>{body}</office:body></office:document-content>'''.encode(), **(assets or {})}
    manifest = f'<manifest:file-entry manifest:full-path="/" manifest:media-type="{mime}"/>'
    for name in entries:
        manifest += f'<manifest:file-entry manifest:full-path="{name}" manifest:media-type="{("image/png" if name.endswith(".png") else "text/xml")}"/>'
    entries['META-INF/manifest.xml'] = (f'''<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">{manifest}</manifest:manifest>''').encode()
    with zipfile.ZipFile(path, 'x', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('mimetype', mime, compress_type=zipfile.ZIP_STORED)
        for name, data in entries.items():
            archive.writestr(name, data)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    root = args.output.resolve()
    root.mkdir(parents=True, exist_ok=False)
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    from PIL import Image, ImageDraw

    note = 'Deployment postponed because the security review found an expired certificate.'
    invoice = 'Invoice 714 for Cedar Labs. Total due 2450 dollars. Payment overdue.'
    doc = Document()
    doc.add_heading('Security review', 0)
    doc.add_paragraph(note)
    doc.add_page_break()
    doc.add_heading('Customer invoice', 0)
    doc.add_paragraph(invoice)
    doc.save(root / 'notes.docx')
    write_odf(root / 'notes.odt', 'text', f'''<office:text>
<text:h text:outline-level="1">Security review</text:h><text:p>{note}</text:p>
<text:h text:outline-level="1" text:style-name="NewPage">Customer invoice</text:h>
<text:p>{invoice}</text:p></office:text>''',
        '<style:style style:name="NewPage" style:family="paragraph"><style:paragraph-properties fo:break-before="page"/></style:style>')
    book = Workbook()
    sheet = book.active
    sheet.title = 'Quarterly sales'
    for row in [['Region', 'Revenue', 'Quarter'], ['North', 42000, 'Q3'], ['South', 17000, 'Q3']]:
        sheet.append(row)
    sheet.column_dimensions['A'].width = 18
    sheet.column_dimensions['B'].width = 18
    book.save(root / 'sales.xlsx')
    rows = ''
    for row in [['Region', 'Revenue', 'Quarter'], ['North', 42000, 'Q3'], ['South', 17000, 'Q3']]:
        rows += '<table:table-row>'
        for value in row:
            attributes = f'office:value-type="float" office:value="{value}"' if isinstance(value, int) else 'office:value-type="string"'
            rows += f'<table:table-cell {attributes}><text:p>{value}</text:p></table:table-cell>'
        rows += '</table:table-row>'
    write_odf(root / 'sales.ods', 'spreadsheet', f'<office:spreadsheet><table:table table:name="Quarterly sales">{rows}</table:table></office:spreadsheet>')

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = 'Quarterly revenue growth'
    data = CategoryChartData()
    data.categories = ['Q1', 'Q2', 'Q3', 'Q4']
    data.add_series('Revenue', [8, 15, 25, 35])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(5), data)
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = 'Security review'
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3)).text = note
    deck.save(root / 'growth.pptx')
    chart = Image.new('RGB', (640, 480), 'white')
    draw = ImageDraw.Draw(chart)
    draw.text((40, 20), 'QUARTERLY REVENUE', fill='black', font_size=30)
    for i, height in enumerate([80, 150, 250, 350]):
        draw.rectangle((80+120*i, 420-height, 160+120*i, 420), fill='blue')
        draw.text((100+120*i, 440), f'Q{i+1}', fill='black', font_size=20)
    import io
    picture = io.BytesIO()
    chart.save(picture, format='PNG')
    write_odf(root / 'growth.odp', 'presentation', f'''<office:presentation>
<draw:page draw:name="Revenue"><draw:frame svg:x="1cm" svg:y="1cm" svg:width="20cm" svg:height="2cm"><draw:text-box><text:p>Quarterly revenue growth</text:p></draw:text-box></draw:frame>
<draw:frame svg:x="5cm" svg:y="3cm" svg:width="14cm" svg:height="10.5cm"><draw:image xlink:href="Pictures/chart.png" xlink:type="simple" xlink:show="embed" xlink:actuate="onLoad"/></draw:frame></draw:page>
<draw:page draw:name="Security"><draw:frame svg:x="1cm" svg:y="1cm" svg:width="20cm" svg:height="10cm"><draw:text-box><text:p>Security review</text:p><text:p>{note}</text:p></draw:text-box></draw:frame></draw:page>
</office:presentation>''', assets={'Pictures/chart.png': picture.getvalue()})
    items = []
    for suffix, family, stem, texts, pages in [
        ('.docx', 'document', 'notes', ['Cedar Labs', 'expired certificate'], 2),
        ('.odt', 'document', 'notes', ['Cedar Labs', 'expired certificate'], 2),
        ('.xlsx', 'spreadsheet', 'sales', ['North', '42000', 'South', '17000'], 1),
        ('.ods', 'spreadsheet', 'sales', ['North', '42000', 'South', '17000'], 1),
        ('.pptx', 'presentation', 'growth', ['Quarterly revenue growth', 'expired certificate'], 2),
        ('.odp', 'presentation', 'growth', ['Quarterly revenue growth', 'expired certificate'], 2)]:
        items.append(dict(id=suffix[1:]+'-layout', modality=family, file=stem+suffix,
                          expectedText=texts, expectedPages=pages,
                          assertChartPixels=suffix in ('.pptx', '.odp')))
    files = {p.name: dict(bytes=p.stat().st_size, sha256=sha256(p)) for p in root.iterdir()}
    (root / 'fixtures.json').write_text(json.dumps(dict(schemaVersion=1,
        purpose='Authored Office layout smoke fixtures, not a quality holdout', items=items, files=files), indent=2)+'\n', encoding='utf-8')
    print(f'Created {len(items)} authored Office fixtures')


if __name__ == '__main__':
    main()
